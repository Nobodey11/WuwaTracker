"""Core features of 0.6 → 1.0: backups, export, import checks, updates, log watcher, planner, wishlist,
what-if, luck tests, Bayesian bands, Prydwen banners/weapons, calendar, rerun tracker, collection, summary,
RAG, benchmark scoring, achievements, data folder."""
import json
import sqlite3
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import numpy as np
import pytest

from wuwa_tracker.core import (backup, benchmark, calendar, collection, import_check, log_watcher, paths, planner,
                               probability as pr, rag, stats_tests, summary, updates)
from wuwa_tracker.core import prydwen as pw
from wuwa_tracker.core.achievements import add_to_board, check as ach_check, leaderboard_entry
from wuwa_tracker.core.analysis import all_stats, default_gamedata
from wuwa_tracker.core.models import Pull
from wuwa_tracker.core.simulator import ConveneSim
from wuwa_tracker.core.storage import Storage

FX = Path(__file__).parent / "fixtures"
UTC = timezone.utc
H = pr.hazard_community()


def mk(pt, name, q, t, g=1, rid=1000, player="1"):
    return Pull(pt, rid, q, name, t, g, player)


def some_pulls():
    t0 = datetime(2026, 10, 1, tzinfo=UTC)
    ps = []
    for i in range(70):
        ps.append(mk(1, "Hsin" if i == 69 else ("Lumi" if i % 10 == 9 else "Sword of Night"),
                     5 if i == 69 else (4 if i % 10 == 9 else 3), t0 + timedelta(hours=i), rid=1311 if i == 69 else 21010013))
    ps.append(mk(2, "Kumokiri", 5, t0 + timedelta(days=5), rid=21020076))
    return ps


# ---------------------------------------------------------------- 0.6
def test_backup_restore_roundtrip(tmp_path):
    st = Storage(tmp_path / "pulls.db")
    st.add_pulls(some_pulls())
    b = st.backup("before-import")
    st.clear()
    assert st.count() == 0
    lst = backup.list_backups(st.path)
    assert lst and lst[0].pulls == 71 and "before-import" in lst[0].reason
    st.restore(lst[0].path)
    assert st.count() == 71
    for i in range(25):
        st.backup(f"n{i}")
    assert len(backup.list_backups(st.path)) == backup.KEEP


def test_exports_roundtrip_json_csv_xlsx(tmp_path):
    st = Storage(tmp_path / "a.db")
    st.add_pulls(some_pulls())
    st.export_wuwatracker_json(tmp_path / "x.json", "1")
    st2 = Storage(tmp_path / "b.db")
    new, total, player = st2.import_wuwatracker_json(tmp_path / "x.json")
    assert (new, total, player) == (71, 71, "1")
    assert [p.key for p in st2.load()] == [p.key for p in st.load()]
    st.export_csv(tmp_path / "x.csv")
    assert (tmp_path / "x.csv").read_text().count("\n") == 72
    pytest.importorskip("openpyxl")
    st.export_xlsx(tmp_path / "x.xlsx", gd=default_gamedata())
    from openpyxl import load_workbook
    wb = load_workbook(tmp_path / "x.xlsx")
    assert wb.sheetnames == ["Pulls", "Banners", "5★ history"]
    assert wb["Pulls"].max_row == 72 and wb["5★ history"].max_row == 3


def test_import_checks():
    old = some_pulls()
    t = datetime(2026, 12, 1, tzinfo=UTC)
    new = [mk(1, "Lumi", 4, t, 1), mk(1, "Sword of Night", 3, t, 3)]        # gap + broken multi-pull
    f = import_check.check_new_vs_existing(old, new)
    assert len(f) == 1 and "Pulls in between may be missing" in f[0].text
    assert import_check.check_ten_pulls(new)[0].level == "warning"
    assert not import_check.check_ten_pulls(old)
    cmp = import_check.compare_with_export(old[:-1], old)
    assert any(x.level == "warning" and x.banner == 2 for x in cmp)
    assert any(x.level == "info" and x.banner == 1 for x in cmp)


def test_update_check():
    assert updates.is_newer("v1.0.0", "0.9.0") and not updates.is_newer("v0.6.0", "0.6.0")
    ok, rel = updates.check("0.6.0", fetch=lambda: {"tag": "v1.0.0", "name": "1.0", "url": "u", "notes": "", "assets": []})
    assert ok and rel["tag"] == "v1.0.0"
    ok, msg = updates.check("0.6.0", fetch=lambda: (_ for _ in ()).throw(OSError("offline")))
    assert not ok and "offline" in msg


# ---------------------------------------------------------------- 0.7
def test_log_watcher_detects_only_new_links(tmp_path):
    from tests.test_game_import import URL, encode_client_log
    f = tmp_path / "Client.log"
    f.write_bytes(encode_client_log(f"old {URL}\n".encode()))
    w = log_watcher.LogWatcher([f])
    assert w.poll() is None                        # priming: existing link is not "new"
    assert w.poll() is None
    new = URL.replace("abc123", "fresh")
    with open(f, "ab") as fh:
        fh.write(encode_client_log(f"noise\n{new}\n".encode()))
    assert w.poll() == new
    assert w.poll() is None
    f.write_bytes(encode_client_log(b"new session\n"))            # truncated
    assert w.poll() is None


def test_data_folder_and_lock(tmp_path, monkeypatch):
    monkeypatch.setenv("APPDATA", str(tmp_path / "appdata"))
    st = Storage(paths.data_dir() / "pulls.db")
    st.add_pulls(some_pulls()); st.con.close()
    target = paths.set_data_dir(tmp_path / "OneDrive" / "Wuwa")
    assert paths.data_dir() == target and (target / "pulls.db").exists()
    assert paths.acquire_lock(target) is None
    (target / "pulls.db.lock").write_text(json.dumps({"holder": "OTHERPC:123", "time": __import__("time").time()}))
    assert "OTHERPC" in paths.acquire_lock(target)
    paths.set_data_dir(None)
    assert paths.data_dir() == paths.app_dir()


def test_accounts_and_income(tmp_path):
    st = Storage(tmp_path / "a.db")
    st.set_nickname("606829202", "Main")
    assert st.nicknames() == {"606829202": "Main"}
    for d in range(10):
        st.add_income((date(2026, 10, 1) + timedelta(days=d)).isoformat(), 100, tides=1 if d == 0 else 0)
    rid = st.add_income("2026-10-10", 1000, source="event")
    rate = planner.income_rate(st.income(), today=date(2026, 10, 10))
    assert rate["days"] == 10 and rate["astrite_per_day"] == pytest.approx(200)
    st.delete_income(rid)
    assert len(st.income()) == 10
    assert planner.income_by_month(st.income())["2026-10"]["astrite"] == 1000


# ---------------------------------------------------------------- 0.8
def test_plan_matches_engine():
    p = planner.plan(H, 8000, 5, date(2026, 11, 11), 100, today=date(2026, 10, 27))
    assert p.days == 15 and p.pulls == (8000 + 1500) // 160 + 5
    assert p.probability == pytest.approx(pr.featured_distribution(H, p.pulls, 1).probability)
    d = planner.date_for_confidence(H, 0, 0, 160, 0.5, today=date(2026, 1, 1))
    assert d == date(2026, 1, 1) + timedelta(days=pr.pulls_needed(H, 1, 0.5))


def test_wishlist_single_equals_engine_and_order():
    w = planner.wishlist(H, H, [planner.Wish("A", "character", 1)], 100, tides=10)
    assert w.all == pytest.approx(pr.featured_distribution(H, 110, 1).probability, abs=1e-9)
    two = planner.wishlist(H, H, [planner.Wish("A"), planner.Wish("W", "weapon")], 150)
    assert two.cumulative[1] <= two.cumulative[0] + 1e-12
    # weapon alone with 80 pulls is certain (hard pity, always featured)
    assert planner.wishlist(H, H, [planner.Wish("W", "weapon")], 80).all == pytest.approx(1.0)
    # forging tides cannot pay for the character
    a = planner.wishlist(H, H, [planner.Wish("A")], 0, forging=200).all
    assert a == pytest.approx(0.0)
    assert "Dropping" in two.advice() or two.all >= 0.9


def test_wishlist_mc_crosscheck():
    """Monte Carlo of pulling sequentially agrees with the exact convolution."""
    import random
    wishes = [planner.Wish("A", "character", 1), planner.Wish("W", "weapon", 1), planner.Wish("B", "character", 1)]
    exact = planner.wishlist(H, H, wishes, 200, tides=20, char_pity=30).cumulative[-1]
    rng = random.Random(1)

    def cost(pity, g, fc):
        n = 0
        while True:
            n += 1; pity += 1
            if rng.random() < H[min(pity, 80)]:
                if g or rng.random() < fc:
                    return n
                g, pity = True, 0
    ok, N = 0, 20000
    for _ in range(N):
        c = cost(30, False, 0.5) + cost(0, False, 0.5)
        wcost = cost(0, False, 1.0)
        ok += max(0, c - 20) + wcost <= 200
    assert ok / N == pytest.approx(exact, abs=0.015)


def test_what_if():
    real = ["won", "lost", "guaranteed", "won"]
    same = planner.what_if_5050(real, {})
    assert same["whatif_featured"] == pytest.approx(3) and same["difference"] == pytest.approx(0)
    flip = planner.what_if_5050(real, {1: "won"})           # won instead of lost -> next one is a coin flip
    assert flip["whatif_featured"] == pytest.approx(1 + 1 + 0.5 + 1)
    avg = planner.what_if_average_luck(H, 10, 600)
    assert avg["expected_pulls"] > 600 and avg["saved_astrite"] > 0


def test_luck_test_and_bands():
    pmf = pr.pity_pmf(H, 0)
    rng = np.random.default_rng(3)
    normal = list(rng.choice(np.arange(len(pmf)), size=40, p=pmf / pmf.sum()))
    t = stats_tests.luck_test(normal, H, sims=1500)
    assert t.ks_p > 0.05 and "Normal" in t.verdict
    bad = [80] * 20
    t2 = stats_tests.luck_test(bad, H, sims=1500)
    assert t2.ks_p < 0.01 and "Unusual" in t2.verdict and "unluckier" in t2.verdict
    assert stats_tests.luck_test([], H) is None
    b = stats_tests.bayes_bands(normal, [10], H, draws=150)
    assert np.all(b.low <= b.high + 1e-12) and b.rate_low <= b.rate_mid <= b.rate_high
    lo, mid, hi = stats_tests.probability_interval(b, 80, max_draws=40)
    assert lo <= mid <= hi


# ---------------------------------------------------------------- 0.9
def test_prydwen_banners_weapons_and_history(tmp_path):
    p = pw.Prydwen(tmp_path)
    p._get = lambda url: (FX / ("prydwen_banners.html" if "banners" in url else "prydwen_weapons.html")).read_bytes()
    bs = p.banners()
    assert [b["name"] for b in bs] == ["Suoming", "Lynae", "Freeze Frame", "Lily"]
    assert bs[1]["rerun"] and bs[3]["section"] == "upcoming" and bs[2]["category"] == "weapon"
    assert pw.banner_times(bs[0], "eu")[0].startswith("2026-10-22T09")
    hist = p.banner_history()
    assert {h["name"] for h in hist} == {"Suoming", "Lynae", "Freeze Frame"}        # Lily has no date yet
    wl = p.weapon_list()
    assert wl["overture"]["rarity"] == 4 and wl["abysssurges"]["resource_id"] == 21040015


def test_tier_history(tmp_path):
    p = pw.Prydwen(tmp_path)
    tl = pw.TierList(["T0"], [pw.TierEntry("T0", "DPS", "Hsin", "hsin")], fetched=1.79e9)
    p._record_tier_snapshot(tl)
    p._record_tier_snapshot(tl)                                       # unchanged: no new snapshot
    tl2 = pw.TierList(["T1"], [pw.TierEntry("T1", "DPS", "Hsin", "hsin")], fetched=1.80e9)
    p._record_tier_snapshot(tl2)
    h = p.tier_history("Hsin")
    assert list(h.values()) == ["T0", "T1"]


def test_calendar_and_reruns():
    gd = default_gamedata()
    bs = pw.parse_banners((FX / "prydwen_banners.html").read_text())
    now = datetime(2026, 11, 9, 12, tzinfo=UTC)
    ev = calendar.build(gd.raw, bs, [{"name": "My event", "start": "2026-11-01", "end": "2026-11-20"}], "Europe", 1, now)
    titles = [e.title for e in ev]
    assert "Suoming" in titles and "Daily reset" in titles and "My event" in titles
    assert titles.count("Suoming") == 1                         # Prydwen and game data merged
    soon = calendar.ending_soon(ev, now)
    assert {e.title for e in soon} >= {"Freeze Frame", "Suoming"}
    pulls = [mk(1, "Hsin", 5, datetime(2026, 10, 1, tzinfo=UTC))]
    runs = {r.name: r for r in calendar.rerun_tracker(sorted(gd.raw["limited_5star_resonators"]), gd.raw, bs, [],
                                                       pulls, today=date(2026, 10, 27))}
    assert runs["Suoming"].live and runs["Hsin"].days_since(date(2026, 10, 27)) == 5
    assert len(runs["Hsin"].runs) == 1                         # pull merged into the game-data run


def test_collection():
    gd = default_gamedata()
    chars = pw.parse_character_list((FX / "prydwen_characters.html").read_text())
    weps = pw.parse_weapon_list((FX / "prydwen_weapons.html").read_text())
    ps = some_pulls() + [mk(1, "Hsin", 5, datetime(2026, 10, 9, tzinfo=UTC), rid=1311)]
    items = collection.build(ps, {"Rebecca": 1}, chars, weps, gd)
    by = {i.name: i for i in items}
    assert by["Hsin"].level == "S1" and by["Kumokiri"].level == "R1" and by["Rebecca"].owned
    assert not by["Ages of Harvest"].owned and by["Overture"].kind == "weapon"
    s = collection.summary(items)
    assert s["resonator_5"][0] >= 1


# ---------------------------------------------------------------- 1.0
def test_weekly_summary():
    gd = default_gamedata()
    ps = some_pulls()
    now = datetime(2026, 10, 6, tzinfo=UTC)
    f = summary.week_facts(ps, all_stats(ps, gd), gd, now)
    assert f["pulls"] == 71 and f["five_stars"][0]["name"] == "Hsin" and f["five_stars"][0]["pity"] == 70
    txt = summary.week_text(f)
    assert "71 pulls" in txt and "Hsin" in txt


def test_rag_search(tmp_path):
    (tmp_path / "notes").mkdir()
    (tmp_path / "notes" / "teams.md").write_text("Lynae team: Lynae, Mornye and Shorekeeper.\n\nOther stuff.")
    (tmp_path / "pr").mkdir()
    (tmp_path / "pr" / "char_hsin.json").write_text(json.dumps({"name": "Hsin", "url": "u", "element": "Havoc", "rarity": 5,
        "roles": [{"role": "DPS", "ratings": {"ToA": "T0"}}], "weapons": [{"name": "Blooming Jadehaven"}], "echo_sets": []}))
    hits = rag.search("best weapon for Hsin", tmp_path / "notes", tmp_path / "pr", default_gamedata().raw)
    assert hits[0]["title"] == "Hsin" and "Blooming Jadehaven" in hits[0]["text"]
    assert rag.search("Lynae team", tmp_path / "notes", None, None)[0]["source"] == "note: teams.md"


def test_benchmark_scoring_and_tools():
    from wuwa_tracker.core.ollama_client import ToolBox, system_prompt
    gd = default_gamedata()
    ps = some_pulls()
    stats = all_stats(ps, gd)
    tb = ToolBox(lambda: stats, lambda: H, lambda: gd)
    qs = benchmark.build_questions(tb)
    assert len(qs) == 6
    q = qs[0]
    assert benchmark.score_answer(f"Your chance is **{q.expected:.1f}%**.", q)
    assert not benchmark.score_answer("about 1 in 3", q)
    assert benchmark.numbers("1,311 pulls and 9.600 Astrite, 46,67 %") == [1311, 9600, 46.67]

    class FakeClient:
        def chat(self, model, msgs, toolbox, on_tool=None, temperature=0.3):
            n = "featured_probability"
            r = toolbox.run(n, {"pool_type": 1, "pulls": 40})
            on_tool(n, {}, r)
            return f"{r['probability_percent']} %", msgs
    res = benchmark.run(FakeClient(), ["m"], tb, qs[:1])
    assert res[0].accuracy == 1 and res[0].tool_rate == 1
    assert "| m | 100% |" in benchmark.report(res)
    tb.income = lambda: {"astrite_per_day": 100}
    out = tb.run("plan_astrite", {"astrite": 8000, "target_date": "2026-12-01"})
    assert out["pulls_available"] >= 50 and "probability_percent" in out
    out = tb.run("wishlist_odds", {"wishes": [{"name": "Lynae"}, {"name": "Spectrum Blaster", "kind": "weapon"}],
                                   "astrite": 32000})
    assert len(out["wishes"]) == 2 and 0 <= out["p_everything_percent"] <= 100
    assert "Hype" not in system_prompt("analyst") and "excited" in system_prompt("hype")


def test_achievements_and_board():
    sim = ConveneSim(H, default_gamedata().rates, False, "Hsin", ["Verina"], pity5=79, guaranteed=True, seed=1)
    batch = sim.pull(10)
    got = ach_check(batch, sim)
    assert "first_gold" in got
    e = leaderboard_entry(sim, "s1")
    board = add_to_board([], e)
    assert board[0]["per_featured"] == e["per_featured"]
    assert add_to_board(board, None) == board


def test_rover_always_owned_and_manual_items():
    gd = default_gamedata()
    items = {i.name: i for i in collection.build([], {}, {}, {}, gd)}
    assert all(items[r].owned and items[r].level == "S0" for r in collection.ROVERS)
    chars = {"roverspectro": {"name": "Rover (Spectro)", "slug": "rover-spectro", "rarity": 5, "element": "Spectro"}}
    items = {i.name: i for i in collection.build([], {"Rover (Spectro)": 2, "Hsin": 1, "Some Gift Weapon": 1}, chars, {}, gd)}
    assert items["Rover (Spectro)"].level == "S2" and "Rover (Havoc)" not in items     # Prydwen list wins
    assert items["Hsin"].owned and items["Hsin"].rarity == 5
    assert items["Some Gift Weapon"].rarity == 4 and items["Some Gift Weapon"].owned
